import json
from pathlib import Path

import docx
import fitz
import pytest

from resume_platform.ingestion.service import DocumentIngestionService


def test_ingest_text_file_persists_raw_text(tmp_path: Path) -> None:
    service = DocumentIngestionService(storage_dir=tmp_path)

    response = service.ingest_bytes(
        filename="resume.txt",
        content=b"Jane Doe\nPython Developer\nFastAPI",
    )

    assert response.filename == "resume.txt"
    assert response.content_type == "txt"
    assert response.character_count == len("Jane Doe\nPython Developer\nFastAPI")
    stored = json.loads(Path(response.storage_path).read_text(encoding="utf-8"))
    assert stored["extracted_text"] == "Jane Doe\nPython Developer\nFastAPI"


def test_ingest_docx_file_extracts_text(tmp_path: Path) -> None:
    service = DocumentIngestionService(storage_dir=tmp_path)
    document = docx.Document()
    document.add_paragraph("John Smith")
    document.add_paragraph("Backend Engineer")

    buffer = Path(tmp_path / "resume.docx")
    document.save(buffer)

    response = service.ingest_bytes(
        filename="resume.docx",
        content=buffer.read_bytes(),
    )

    stored = json.loads(Path(response.storage_path).read_text(encoding="utf-8"))
    assert "John Smith" in stored["extracted_text"]
    assert "Backend Engineer" in stored["extracted_text"]


def test_ingest_pdf_file_extracts_text(tmp_path: Path) -> None:
    service = DocumentIngestionService(storage_dir=tmp_path)
    pdf_path = tmp_path / "resume.pdf"

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Alice Johnson Resume")
    document.save(pdf_path)
    document.close()

    response = service.ingest_bytes(
        filename="resume.pdf",
        content=pdf_path.read_bytes(),
    )

    stored = json.loads(Path(response.storage_path).read_text(encoding="utf-8"))
    assert "Alice Johnson Resume" in stored["extracted_text"]


def test_rejects_unsupported_file_types(tmp_path: Path) -> None:
    service = DocumentIngestionService(storage_dir=tmp_path)

    with pytest.raises(ValueError, match="Unsupported file type"):
        service.ingest_bytes(filename="resume.csv", content=b"name,skills")


def test_rejects_empty_upload(tmp_path: Path) -> None:
    service = DocumentIngestionService(storage_dir=tmp_path)

    with pytest.raises(ValueError, match="Uploaded file is empty"):
        service.ingest_bytes(filename="resume.txt", content=b"")
