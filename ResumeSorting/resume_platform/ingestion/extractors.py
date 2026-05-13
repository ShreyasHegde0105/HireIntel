from __future__ import annotations

from abc import ABC, abstractmethod
from io import BytesIO

import docx
import fitz
import pdfplumber


class BaseExtractor(ABC):
    @abstractmethod
    def extract(self, content: bytes) -> str:
        raise NotImplementedError


class PdfExtractor(BaseExtractor):
    def extract(self, content: bytes) -> str:
        text = self._extract_with_pymupdf(content)
        if text.strip():
            return text

        text = self._extract_with_pdfplumber(content)
        if text.strip():
            return text

        raise ValueError("No readable text could be extracted from the PDF file.")

    @staticmethod
    def _extract_with_pymupdf(content: bytes) -> str:
        with fitz.open(stream=content, filetype="pdf") as document:
            pages = [page.get_text("text") for page in document]
        return _normalize_text("\n".join(pages))

    @staticmethod
    def _extract_with_pdfplumber(content: bytes) -> str:
        with pdfplumber.open(BytesIO(content)) as document:
            pages = [page.extract_text() or "" for page in document.pages]
        return _normalize_text("\n".join(pages))


class DocxExtractor(BaseExtractor):
    def extract(self, content: bytes) -> str:
        document = docx.Document(BytesIO(content))
        chunks: list[str] = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    chunks.append(row_text)

        text = _normalize_text("\n".join(chunks))
        if not text:
            raise ValueError("No readable text could be extracted from the DOCX file.")
        return text


class TextExtractor(BaseExtractor):
    def extract(self, content: bytes) -> str:
        try:
            decoded = content.decode("utf-8")
        except UnicodeDecodeError:
            decoded = content.decode("latin-1")
        text = _normalize_text(decoded)
        if not text:
            raise ValueError("No readable text could be extracted from the text file.")
        return text


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in text.splitlines()]
    filtered = [line for line in lines if line]
    return "\n".join(filtered).strip()
