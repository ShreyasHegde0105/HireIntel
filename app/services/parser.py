import io
import logging
import re
import zipfile

from docx import Document

logger = logging.getLogger(__name__)


class DocumentError(ValueError):
    """A problem with an uploaded document. The message is safe to show to users."""


IMAGE_MIME_TYPES = {"png": "image/png", "jpeg": "image/jpeg", "webp": "image/webp"}
EXTENSION_TO_FORMAT = {
    "pdf": "pdf",
    "docx": "docx",
    "txt": "txt",
    "png": "png",
    "jpg": "jpeg",
    "jpeg": "jpeg",
    "webp": "webp",
}
SUPPORTED_HINT = "Supported formats: PDF, DOCX, TXT, PNG, JPG, WEBP."


# pypdfium2 marks a hyphenated line break with U+FFFE after joining the word halves
# ("soft￾ware"), and Word uses U+00AD soft hyphens; both are dropped. Other control
# characters (often from icon fonts in contact lines) are dropped too, keeping newlines and tabs.
_INVISIBLE_CHARS = re.compile("[￾­\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")


def clean_text(text: str) -> str:
    """Normalises extracted text so words match the way they read in the document."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return _INVISIBLE_CHARS.sub("", text).strip()


def _extension(filename: str) -> str:
    return filename.lower().rsplit(".", 1)[-1] if "." in filename else ""


def _sniff(file_bytes: bytes) -> str:
    """Identifies the real file type from its first bytes. Returns '' when unknown."""
    head = file_bytes[:1024]
    if b"%PDF-" in head:
        return "pdf"
    if file_bytes.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"  # legacy .doc / .xls / .ppt
    if file_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if file_bytes.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if file_bytes[:4] == b"RIFF" and file_bytes[8:12] == b"WEBP":
        return "webp"
    if file_bytes.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                if "word/document.xml" in zf.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
        return "zip"
    return ""


def detect_format(file_bytes: bytes, filename: str) -> str:
    """Checks that the extension is supported and matches the file's actual content.

    Returns one of: pdf, docx, txt, png, jpeg, webp.
    """
    if not file_bytes:
        raise DocumentError("The file is empty.")

    ext = _extension(filename)
    if ext == "doc":
        raise DocumentError("Legacy .doc files are not supported. Please re-save the file as DOCX or PDF.")
    expected = EXTENSION_TO_FORMAT.get(ext)
    if expected is None:
        if not ext:
            raise DocumentError(f"The file has no extension. {SUPPORTED_HINT}")
        raise DocumentError(f"Unsupported file type '.{ext}'. {SUPPORTED_HINT}")

    actual = _sniff(file_bytes)
    if expected == "txt":
        if actual or b"\x00" in file_bytes[:4096]:
            raise DocumentError("This .txt file does not contain plain text.")
        return "txt"
    if actual == "ole":
        raise DocumentError("This looks like a legacy Word (.doc) file. Please re-save it as DOCX or PDF.")
    if actual != expected:
        raise DocumentError(f"The file content does not match its .{ext} extension.")
    return expected


def parse_pdf(file_bytes: bytes) -> str:
    """Extracts the text layer of a PDF.

    Returns '' for a readable PDF with no text layer (a scan), so the caller can try OCR.
    Raises DocumentError for corrupt or password-protected files, which OCR cannot fix.
    """
    errors = []

    # 1. Primary fast C++ engine: pypdfium2
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(file_bytes)
        try:
            extracted = []
            for page in pdf:
                textpage = page.get_textpage()
                page_text = textpage.get_text_range()
                if page_text:
                    extracted.append(page_text)
                textpage.close()
                page.close()
        finally:
            pdf.close()
        text = "\n".join(extracted).strip()
        if text:
            return text
    except Exception as e:
        errors.append(e)

    # 2. Secondary: pdfplumber
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages_text = [pt for pt in (page.extract_text() for page in pdf.pages) if pt]
        return "\n".join(pages_text).strip()
    except Exception as e:
        errors.append(e)

    logger.info("pdf could not be opened", extra={"errors": [type(e).__name__ for e in errors]})
    raise DocumentError("This PDF could not be opened. It may be corrupt or password-protected.")


def parse_docx(file_bytes: bytes) -> str:
    """Extracts text from a DOCX file: body paragraphs and tables.

    Known gap: text boxes are not read yet.
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception:
        raise DocumentError("This DOCX file could not be opened. It may be corrupt.") from None

    text_content = []
    for paragraph in doc.paragraphs:
        clean_p = paragraph.text.strip()
        if clean_p:
            text_content.append(clean_p)

    # Tables are often used for resume layout
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text and cell_text not in row_text:
                    row_text.append(cell_text)
            if row_text:
                text_content.append(" | ".join(row_text))

    return "\n".join(text_content).strip()


def extract_text(file_bytes: bytes, file_format: str) -> str:
    """Extracts raw text in memory. Returns '' for images and scanned PDFs (OCR is the caller's job)."""
    if file_format == "pdf":
        return parse_pdf(file_bytes)
    if file_format == "docx":
        return parse_docx(file_bytes)
    if file_format == "txt":
        return file_bytes.decode("utf-8", errors="replace").strip()
    if file_format in IMAGE_MIME_TYPES:
        return ""
    raise DocumentError(f"Unsupported file type. {SUPPORTED_HINT}")
